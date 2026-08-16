from io import BytesIO

from openpyxl import Workbook

from thermovisi.excel_parser import parse_workbook, preview_rows


def workbook_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "BAY1"
    ws["B15"] = "Titik A"
    ws["J15"] = 40
    ws["K15"] = "41,5"
    ws["L15"] = 42
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_parse_phase_and_delta():
    items = [{
        "template_item_id": 1, "form_section_code": "MAIN", "sequence_no": 1,
        "source_sheet_pattern": r"^BAY\d+$", "source_row_no": 15, "source_occurrence_no": 1,
        "raw_point_label": "Titik A", "equipment_group_code": "PMT", "point_code": "POINT_A",
        "measurement_mode_code": "PHASE", "phase_codes": ["R", "S", "T"],
        "source_value_map": {"R": "J", "S": "K", "T": "L"}, "is_required": True,
    }]
    sheets, warnings = parse_workbook(workbook_bytes(), items, 30.0)
    assert warnings == []
    assert sheets[0].numeric_count == 3
    assert [m.temperature_c for m in sheets[0].measurements] == [40.0, 41.5, 42.0]
    assert [m.delta_ambient_c for m in sheets[0].measurements] == [10.0, 11.5, 12.0]


def test_label_search_single_value():
    wb = Workbook()
    ws = wb.active
    ws.title = "A. CUSTOM"
    ws["D20"] = "Suhu NGR Atas"
    ws["E20"] = 37.25
    buffer = BytesIO()
    wb.save(buffer)
    items = [{
        "template_item_id": 9, "form_section_code": "A", "sequence_no": 1,
        "source_sheet_pattern": r"^A\.", "source_row_no": None, "source_occurrence_no": 1,
        "raw_point_label": "Suhu NGR Atas", "equipment_group_code": "NGR", "point_code": "NGR_TOP",
        "measurement_mode_code": "SINGLE", "phase_codes": [],
        "source_value_map": {"VALUE": "E"}, "is_required": True,
    }]
    sheets, _ = parse_workbook(buffer.getvalue(), items, 30.0)
    assert sheets[0].measurements[0].source_cell_address == "E20"
    assert sheets[0].measurements[0].delta_ambient_c == 7.25


def test_preview_uses_ambient_for_each_sheet():
    items = [{
        "template_item_id": 1, "form_section_code": "MAIN", "sequence_no": 1,
        "source_sheet_pattern": r"^BAY\d+$", "source_row_no": 15, "source_occurrence_no": 1,
        "raw_point_label": "Titik A", "equipment_group_code": "PMT", "point_code": "POINT_A",
        "measurement_mode_code": "PHASE", "phase_codes": ["R", "S", "T"],
        "source_value_map": {"R": "J", "S": "K", "T": "L"}, "is_required": True,
    }]
    sheets, _ = parse_workbook(workbook_bytes(), items)
    rows = preview_rows(sheets, {"BAY1": 35.0})
    assert [row["delta_ambient_c"] for row in rows] == [5.0, 6.5, 7.0]


def test_sheet_name_does_not_have_to_match_template_pattern():
    items = [{
        "template_item_id": 1, "form_section_code": "MAIN", "sequence_no": 1,
        "source_sheet_pattern": r"^NAMA_LAMA_YANG_TIDAK_COCOK$", "source_row_no": 15,
        "source_occurrence_no": 1, "raw_point_label": "Titik A",
        "equipment_group_code": "PMT", "point_code": "POINT_A",
        "measurement_mode_code": "PHASE", "phase_codes": ["R", "S", "T"],
        "source_value_map": {"R": "J", "S": "K", "T": "L"}, "is_required": True,
    }]

    sheets, warnings = parse_workbook(workbook_bytes(), items, 30.0)

    assert warnings == []
    assert [sheet.sheet_name for sheet in sheets] == ["BAY1"]
    assert sheets[0].numeric_count == 3


def test_form_section_is_detected_from_labels_not_sheet_name():
    wb = Workbook()
    ws = wb.active
    ws.title = "Nama Bebas Unit Toba"
    ws["B10"] = "Titik Trafo"
    ws["E10"] = 45
    buffer = BytesIO()
    wb.save(buffer)
    common = {
        "source_sheet_pattern": r"TIDAK_DIPAKAI", "source_row_no": None,
        "source_occurrence_no": 1, "equipment_group_code": "TRF",
        "measurement_mode_code": "SINGLE", "phase_codes": [],
        "source_value_map": {"VALUE": "E"}, "is_required": True,
    }
    items = [
        dict(common, template_item_id=1, form_section_code="A", sequence_no=1,
             raw_point_label="Titik Trafo", point_code="TRF_POINT"),
        dict(common, template_item_id=2, form_section_code="B", sequence_no=1,
             raw_point_label="Titik Bay", point_code="BAY_POINT"),
    ]

    sheets, warnings = parse_workbook(buffer.getvalue(), items)

    assert warnings == []
    assert sheets[0].measurements[0].form_section_code == "A"
    assert sheets[0].measurements[0].temperature_c == 45


def test_dash_is_not_measured_and_does_not_make_sheet_invalid():
    wb = Workbook()
    ws = wb.active
    ws.title = "BAY TIDAK DIUKUR"
    ws["B15"] = "Titik A"
    ws["J15"] = "-"
    ws["K15"] = "—"
    ws["L15"] = "TIDAK DIUKUR"
    buffer = BytesIO()
    wb.save(buffer)
    items = [{
        "template_item_id": 1, "form_section_code": "MAIN", "sequence_no": 1,
        "source_row_no": 15, "source_occurrence_no": 1, "raw_point_label": "Titik A",
        "equipment_group_code": "PMT", "point_code": "POINT_A",
        "measurement_mode_code": "PHASE", "phase_codes": ["R", "S", "T"],
        "source_value_map": {"R": "J", "S": "K", "T": "L"}, "is_required": True,
    }]

    sheets, warnings = parse_workbook(buffer.getvalue(), items)

    assert warnings == []
    assert len(sheets) == 1
    assert sheets[0].numeric_count == 0
    assert sheets[0].invalid_count == 0
    assert sheets[0].not_measured_count == 3
    assert {m.data_quality_status for m in sheets[0].measurements} == {"NOT_MEASURED"}


def test_blank_required_cell_remains_invalid():
    wb = Workbook()
    ws = wb.active
    ws["B15"] = "Titik A"
    buffer = BytesIO()
    wb.save(buffer)
    items = [{
        "template_item_id": 1, "form_section_code": "MAIN", "sequence_no": 1,
        "source_row_no": 15, "source_occurrence_no": 1, "raw_point_label": "Titik A",
        "equipment_group_code": "PMT", "point_code": "POINT_A",
        "measurement_mode_code": "SINGLE", "phase_codes": [],
        "source_value_map": {"VALUE": "E"}, "is_required": True,
    }]
    sheets, warnings = parse_workbook(buffer.getvalue(), items)
    assert sheets == []
    assert any("tidak berisi nilai angka" in warning for warning in warnings)
