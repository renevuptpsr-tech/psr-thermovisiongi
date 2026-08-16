from __future__ import annotations

import io
import math
import re
from dataclasses import asdict, dataclass
from typing import Any

from openpyxl import load_workbook


@dataclass(slots=True)
class ParsedMeasurement:
    sheet_name: str
    form_section_code: str
    sequence_no: int
    template_item_id: int
    equipment_group_code: str
    point_code: str
    point_label: str
    phase_code: str | None
    temperature_c: float | None
    delta_ambient_c: float | None
    source_cell_address: str | None
    source_value_raw: str | None
    data_quality_status: str
    validation_message: str | None

    def as_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass(slots=True)
class ParsedSheet:
    sheet_name: str
    measurements: list[ParsedMeasurement]
    numeric_count: int
    invalid_count: int
    warning_count: int
    not_measured_count: int = 0
    not_applicable_count: int = 0


def _normalise(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip().casefold()


def _number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = str(value).strip().replace("°C", "").replace("°", "")
        text = text.replace(" ", "").replace(",", ".")
        try:
            number = float(text)
        except ValueError:
            return None
    return number if math.isfinite(number) else None


def _row_for_label(ws, label: str, occurrence: int) -> int | None:
    needle = _normalise(label)
    if not needle:
        return None
    hits = 0
    # Label form berada di sisi kiri; batasi agar angka/hasil tidak ikut dipindai.
    for row in ws.iter_rows(min_col=1, max_col=min(ws.max_column, 9)):
        for cell in row:
            if _normalise(cell.value) == needle:
                hits += 1
                if hits == occurrence:
                    return cell.row
    return None


def _sheet_label_rows(ws) -> dict[str, list[int]]:
    """Indeks label sisi kiri sheet; nama sheet sengaja tidak digunakan."""
    rows_by_label: dict[str, list[int]] = {}
    for row in ws.iter_rows(min_col=1, max_col=min(ws.max_column, 9)):
        for cell in row:
            label = _normalise(cell.value)
            if label:
                rows_by_label.setdefault(label, []).append(cell.row)
    return rows_by_label


def _items_for_sheet(
    ws,
    template_items: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], dict[str, list[int]], str | None]:
    """Pilih bagian template dari label isi sheet, bukan dari nama sheet."""
    label_rows = _sheet_label_rows(ws)
    items_by_section: dict[str, list[dict[str, Any]]] = {}
    for item in template_items:
        section = str(item.get("form_section_code") or "MAIN")
        items_by_section.setdefault(section, []).append(item)

    scores = {
        section: sum(
            1
            for item in items
            if _normalise(item.get("raw_point_label")) in label_rows
        )
        for section, items in items_by_section.items()
    }
    best_score = max(scores.values(), default=0)
    if best_score == 0:
        return [], label_rows, None
    winners = [section for section, score in scores.items() if score == best_score]
    if len(winners) > 1:
        return [], label_rows, (
            "Isi sheet cocok sama kuat dengan beberapa bagian template: " + ", ".join(winners)
        )
    return items_by_section[winners[0]], label_rows, None


def _value_spec(value_map: dict[str, Any], key: str) -> tuple[str | None, int]:
    """Mendukung map sederhana {R: J} maupun {R: {column: J, row_offset: 0}}."""
    spec = value_map.get(key)
    if isinstance(spec, str):
        return spec.strip().upper(), 0
    if isinstance(spec, dict):
        column = spec.get("column") or spec.get("col")
        return (str(column).strip().upper() if column else None), int(spec.get("row_offset", 0))
    return None, 0


_NOT_MEASURED_VALUES = {"-", "—", "–", "tidak diukur", "tidak dilakukan pengukuran", "not measured"}
_NOT_APPLICABLE_VALUES = {"n/a", "na", "tidak ada", "tidak tersedia", "not applicable"}


def _quality(raw: Any, number: float | None, required: bool) -> tuple[str, str | None]:
    raw_text = _normalise(raw)
    if raw_text in _NOT_MEASURED_VALUES:
        return "NOT_MEASURED", "Pengukuran tidak dilakukan"
    if raw_text in _NOT_APPLICABLE_VALUES:
        return "NOT_APPLICABLE", "Titik ukur tidak tersedia/tidak berlaku"
    if raw is None or raw_text == "":
        return (
            ("INVALID", "Nilai suhu wajib kosong")
            if required
            else ("WARNING", "Nilai opsional kosong")
        )
    if number is None:
        return "INVALID", "Nilai suhu bukan angka yang valid"
    if number < -50 or number > 300:
        return "WARNING", "Nilai suhu di luar rentang operasional wajar (-50 s.d. 300 °C)"
    return "VALID", None


def parse_workbook(
    file_bytes: bytes,
    template_items: list[dict[str, Any]],
    ambient_temperature_c: float | None = None,
) -> tuple[list[ParsedSheet], list[str]]:
    if not template_items:
        raise ValueError("Template tidak mempunyai item aktif.")

    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=False)
    warnings: list[str] = []
    result: list[ParsedSheet] = []

    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        selected_items, label_rows, selection_warning = _items_for_sheet(ws, template_items)
        if selection_warning:
            warnings.append(f"Sheet '{sheet_name}' dilewati: {selection_warning}.")
            continue
        if not selected_items:
            continue
        measurements: list[ParsedMeasurement] = []
        for item in selected_items:
            label = _normalise(item.get("raw_point_label"))
            occurrence = int(item.get("source_occurrence_no") or 1)
            matching_rows = label_rows.get(label, [])
            row_no = matching_rows[occurrence - 1] if len(matching_rows) >= occurrence else None
            # Fallback hanya setelah bagian form dikenali dari isi sheet.
            if row_no is None:
                row_no = item.get("source_row_no")

            mode = item["measurement_mode_code"]
            keys = list(item.get("phase_codes") or []) if mode == "PHASE" else ["VALUE"]
            value_map = item.get("source_value_map") or {}
            for key in keys:
                column, row_offset = _value_spec(value_map, key)
                address = f"{column}{row_no + row_offset}" if column and row_no else None
                raw = ws[address].value if address else None
                temperature = _number(raw)
                quality, message = _quality(raw, temperature, bool(item.get("is_required", True)))
                if row_no is None:
                    quality, message = "INVALID", "Baris titik ukur tidak ditemukan"
                elif column is None:
                    quality, message = "INVALID", f"Pemetaan kolom {key} tidak tersedia"
                delta = (
                    round(temperature - ambient_temperature_c, 3)
                    if temperature is not None and ambient_temperature_c is not None
                    else None
                )
                measurements.append(
                    ParsedMeasurement(
                        sheet_name=sheet_name,
                        form_section_code=item["form_section_code"],
                        sequence_no=int(item["sequence_no"]),
                        template_item_id=int(item["template_item_id"]),
                        equipment_group_code=item["equipment_group_code"],
                        point_code=item["point_code"],
                        point_label=item["raw_point_label"],
                        phase_code=key if mode == "PHASE" else None,
                        temperature_c=temperature,
                        delta_ambient_c=delta,
                        source_cell_address=address,
                        source_value_raw=None if raw is None else str(raw),
                        data_quality_status=quality,
                        validation_message=message,
                    )
                )

        numeric_count = sum(m.temperature_c is not None for m in measurements)
        not_measured_count = sum(
            m.data_quality_status == "NOT_MEASURED" for m in measurements
        )
        not_applicable_count = sum(
            m.data_quality_status == "NOT_APPLICABLE" for m in measurements
        )
        if numeric_count == 0 and not_measured_count == 0 and not_applicable_count == 0:
            warnings.append(
                f"Sheet '{sheet_name}' dikenali dari label titik ukur tetapi tidak berisi nilai angka; "
                "sheet dilewati."
            )
            continue
        result.append(
            ParsedSheet(
                sheet_name=sheet_name,
                measurements=measurements,
                numeric_count=numeric_count,
                invalid_count=sum(m.data_quality_status == "INVALID" for m in measurements),
                warning_count=sum(m.data_quality_status == "WARNING" for m in measurements),
                not_measured_count=not_measured_count,
                not_applicable_count=not_applicable_count,
            )
        )

    if not result:
        warnings.append(
            "Tidak ada sheet yang isi titik ukurnya cocok dengan template yang dipilih. "
            "Nama sheet tidak digunakan sebagai syarat."
        )
    return result, warnings


def preview_rows(
    parsed_sheets: list[ParsedSheet],
    ambient_by_sheet: dict[str, float] | None = None,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for parsed in parsed_sheets:
        ambient = ambient_by_sheet.get(parsed.sheet_name) if ambient_by_sheet else None
        for measurement in parsed.measurements:
            row = measurement.as_dict()
            if ambient is not None and measurement.temperature_c is not None:
                row["delta_ambient_c"] = round(measurement.temperature_c - ambient, 3)
            rows.append(row)
    return rows
